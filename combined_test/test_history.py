"""Lightweight, cached lookup for completed tests of the same shell SN."""

from __future__ import annotations

import math
import threading
from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .excel_export import sanitize_sn
from .shipping_report import (
    REPORT_DATA_SHEET,
    WorkbookCompatibility,
    _compatibility,
    _parse_session_information,
    _result_column_key,
)


@dataclass(frozen=True)
class TestHistoryPoint:
    __test__ = False

    current_a: float
    power_w: float


@dataclass(frozen=True)
class TestHistoryInspection:
    """The small subset of workbook data needed by history list/reference UI."""

    __test__ = False

    source_path: Path
    sn: str
    session: dict[str, str]
    points: tuple[TestHistoryPoint, ...]
    compatibility: WorkbookCompatibility


@dataclass(frozen=True)
class TestHistoryRecord:
    __test__ = False

    inspection: TestHistoryInspection
    modified_time_ns: int

    @property
    def station(self) -> str:
        return self.inspection.session.get("测试站别", "").strip()

    @property
    def started_at_utc(self) -> str:
        return self.inspection.session.get("开始时间 UTC", "").strip()

    def summary_payload(self) -> dict[str, Any]:
        currents = [point.current_a for point in self.inspection.points]
        return {
            "sourcePath": str(self.inspection.source_path),
            "station": self.station,
            "startedAtUtc": self.started_at_utc,
            "modifiedAtMs": self.modified_time_ns / 1_000_000,
            "pointCount": len(currents),
            "minimumCurrentA": min(currents),
            "maximumCurrentA": max(currents),
        }

    def reference_payload(self) -> dict[str, Any]:
        return {
            **self.summary_payload(),
            "points": [
                {
                    "currentA": point.current_a,
                    "powerW": point.power_w,
                }
                for point in self.inspection.points
            ],
        }


# A workbook is immutable for indexing purposes until its mtime changes. Cache
# invalid workbooks too so focus refreshes do not repeatedly reopen them. Keep
# the process-wide cache bounded because operators can visit many SN directories
# during one long-running backend session.
HISTORY_INDEX_CACHE_MAX_ENTRIES = 512
_HISTORY_INDEX_CACHE: OrderedDict[
    tuple[str, int], TestHistoryInspection | None
] = OrderedDict()
_HISTORY_INDEX_CACHE_LOCK = threading.RLock()


def clear_test_history_cache() -> None:
    """Clear the process cache; primarily useful for deterministic tests."""

    with _HISTORY_INDEX_CACHE_LOCK:
        _HISTORY_INDEX_CACHE.clear()


def _finite_number(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _load_history_index(source: Path) -> TestHistoryInspection:
    """Read only session fields plus current/power columns, never spectra."""

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{exc}") from exc
    try:
        session = _parse_session_information(workbook)
        schema_version = str(session.get("出货报告格式版本", "")).strip()
        sn = session.get("SN", "").strip()

        sheet = (
            workbook[REPORT_DATA_SHEET]
            if REPORT_DATA_SHEET in workbook.sheetnames
            else workbook[workbook.sheetnames[0]]
        )
        header_row = 2
        columns: dict[str, int] = {}
        scan_row_limit = min(sheet.max_row, 100)
        scan_column_limit = min(sheet.max_column, 32)
        for row_number, values in enumerate(
            sheet.iter_rows(
                min_row=1,
                max_row=scan_row_limit,
                min_col=1,
                max_col=scan_column_limit,
                values_only=True,
            ),
            start=1,
        ):
            candidate: dict[str, int] = {}
            for column_number, value in enumerate(values, start=1):
                key = _result_column_key(value)
                if key in {"current", "power"} and key not in candidate:
                    candidate[key] = column_number
            if {"current", "power"}.issubset(candidate):
                header_row = row_number
                columns = candidate
                break
        if not columns:
            # Fixed-layout fallback used by early exports.
            columns = {"current": 1, "power": 3}

        maximum_column = max(columns.values())
        points: list[TestHistoryPoint] = []
        for values in sheet.iter_rows(
            min_row=header_row + 1,
            min_col=1,
            max_col=maximum_column,
            values_only=True,
        ):
            current = _finite_number(values[columns["current"] - 1])
            power = _finite_number(values[columns["power"] - 1])
            if current is None or power is None:
                continue
            points.append(TestHistoryPoint(current, power))
        if not points:
            raise ValueError("Excel 中没有有效测试点")
        compatibility = _compatibility(
            session,
            schema_version,
            point_count=len(points),
        )
        return TestHistoryInspection(
            source_path=source,
            sn=sn,
            session=session,
            points=tuple(points),
            compatibility=compatibility,
        )
    finally:
        workbook.close()


def _cached_history_index(source: Path, modified_time_ns: int) -> TestHistoryInspection | None:
    path_key = str(source)
    cache_key = (path_key, modified_time_ns)
    with _HISTORY_INDEX_CACHE_LOCK:
        if cache_key in _HISTORY_INDEX_CACHE:
            _HISTORY_INDEX_CACHE.move_to_end(cache_key)
            return _HISTORY_INDEX_CACHE[cache_key]

    try:
        inspection = _load_history_index(source)
    except (OSError, ValueError):
        inspection = None

    with _HISTORY_INDEX_CACHE_LOCK:
        stale_keys = [key for key in _HISTORY_INDEX_CACHE if key[0] == path_key and key != cache_key]
        for stale_key in stale_keys:
            del _HISTORY_INDEX_CACHE[stale_key]
        _HISTORY_INDEX_CACHE[cache_key] = inspection
        _HISTORY_INDEX_CACHE.move_to_end(cache_key)
        while len(_HISTORY_INDEX_CACHE) > HISTORY_INDEX_CACHE_MAX_ENTRIES:
            _HISTORY_INDEX_CACHE.popitem(last=False)
    return inspection


def _record_sort_key(record: TestHistoryRecord) -> tuple[int, int, str]:
    try:
        started_time_ns = int(datetime.fromisoformat(record.started_at_utc).timestamp() * 1_000_000_000)
    except (TypeError, ValueError, OverflowError):
        started_time_ns = 0
    return started_time_ns, record.modified_time_ns, str(record.inspection.source_path)


def find_test_history(output_dir: Path | str, sn: str) -> tuple[TestHistoryRecord, ...]:
    requested_sn = str(sn).strip()
    if not requested_sn:
        return ()
    root = Path(output_dir).expanduser().resolve()
    sn_directory = (root / sanitize_sn(requested_sn)).resolve()
    if not sn_directory.is_dir():
        return ()

    records: list[TestHistoryRecord] = []
    for path in sn_directory.rglob("*.xlsx"):
        if path.is_symlink() or path.name.startswith("~$"):
            continue
        source = path.resolve()
        try:
            source.relative_to(sn_directory)
            modified_time_ns = source.stat().st_mtime_ns
        except OSError:
            continue
        inspection = _cached_history_index(source, modified_time_ns)
        if inspection is None or inspection.sn != requested_sn:
            continue
        if inspection.compatibility.kind != "verified_success":
            continue
        records.append(TestHistoryRecord(inspection, modified_time_ns))

    return tuple(sorted(records, key=_record_sort_key, reverse=True))
