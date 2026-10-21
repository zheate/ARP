from __future__ import annotations

import copy
import json
import math
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from combined_test.automation import AutomaticTestState
from combined_test.excel_export import ExcelTestRecord, save_test_records
from combined_test.models import PowerMeterOption
from openpyxl import load_workbook
from tauri_bridge.headless_backend import HeadlessBackend
from tauri_bridge.settings import JsonSettings


class PowerMeterProfileTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.settings = JsonSettings(Path(self.temp_dir.name) / "settings.json")
        self.backend = HeadlessBackend(self.settings)

    def tearDown(self) -> None:
        self.backend.export_worker.graceful_stop(timeout_s=5.0)
        self.temp_dir.cleanup()

    @staticmethod
    def option(
        serial: str,
        *,
        driver: str = "caihuang",
        resource: str = "ASRL3::INSTR",
    ) -> PowerMeterOption:
        return PowerMeterOption(
            resource=resource,
            device_type="Test power meter",
            detail=f"probe SN {serial}" if serial else "probe SN unavailable",
            driver_kind=driver,
            serial_number=serial,
        )

    def bind(self, option: PowerMeterOption) -> dict:
        if option not in self.backend.power_meter_options:
            self.backend.power_meter_options.append(option)
        return self.backend.dispatch(
            "app.configure", {"stationId": "1", "powerMeterResource": option.resource}
        )

    def command(self, action: str, **params: object) -> dict:
        config = self.backend.runtime.configuration
        return self.backend.dispatch(
            f"powerMeter.profile.{action}",
            {
                "stationId": "1",
                "driverKind": config["powerMeterDriverKind"],
                "serialNumber": config["powerMeterSerialNumber"],
                **params,
            },
        )

    def library(self) -> dict:
        return self.backend.snapshot()["devices"]["powerMeter"]["coefficientLibrary"]

    def selected_profile(self) -> tuple[dict, dict]:
        library = self.library()
        profile = next(
            item
            for item in library["profiles"]
            if item["id"] == library["selectedProfileId"]
        )
        revision = next(
            item
            for item in profile["versions"]
            if item["version"] == library["selectedVersion"]
        )
        return profile, revision

    def test_delete_selected_profile_switches_to_remaining_profile_and_persists(self) -> None:
        self.bind(self.option("PROBE-A"))
        default_id = self.library()["selectedProfileId"]
        self.command("save", profileId=default_id, expectedVersion=1,
                     name="Updated default", softwareGain=1.07512345, wavelengthNm=975.5)
        self.command("create", name="Remove me", softwareGain=1.2, wavelengthNm=980.0)
        deleted_id = self.library()["selectedProfileId"]
        other_probe = copy.deepcopy(self.backend.power_meter_profiles.ensure("coherent", "OTHER", wavelength=976.0))

        result = self.command("delete", profileId=deleted_id, expectedVersion=1)

        self.assertEqual([p["id"] for p in self.library()["profiles"]], [default_id])
        self.assertEqual(result["configuration"]["powerMeterProfileId"], default_id)
        self.assertEqual(result["configuration"]["softwareGain"], 1.07512345)
        self.assertEqual(result["configuration"]["powerMeterWavelengthNm"], 975.5)
        self.assertEqual(result["configuration"]["powerMeterProfileVersion"], 2)
        saved = json.loads(self.settings.path.read_text(encoding="utf-8"))
        self.assertEqual(saved["runtime/deviceProfiles"], self.backend.device_profiles)
        self.assertEqual(saved["runtime/configuration"]["powerMeterProfileId"], default_id)
        self.assertEqual(self.backend.device_profiles["powerMeters"]["coherent:OTHER"], other_probe)

    def test_delete_unselected_profile_keeps_active_measurement_parameters(self) -> None:
        self.bind(self.option("PROBE-A"))
        default_id = self.library()["selectedProfileId"]
        self.command("create", name="Keep me", softwareGain=1.23456789, wavelengthNm=979.123456)
        before = dict(self.backend.runtime.configuration)
        self.command("delete", profileId=default_id, expectedVersion=1)
        self.assertEqual(self.backend.runtime.configuration, before)
        self.assertEqual(len(self.library()["profiles"]), 1)

    def test_delete_last_profile_stays_empty_after_restart_and_blocks_acquisition(self) -> None:
        meter = self.option("PROBE-A")
        self.bind(meter)
        profile_id = self.library()["selectedProfileId"]
        self.command("delete", profileId=profile_id, expectedVersion=1)
        self.assertEqual(self.library()["profiles"], [])
        self.assertEqual(self.library()["selectedProfileId"], "")
        self.assertEqual(self.library()["selectedVersion"], 0)
        self.assertEqual(self.backend.runtime.configuration["powerMeterProfileName"], "未绑定方案")
        with self.assertRaisesRegex(ValueError, "建立系数方案"):
            self.backend.runtime.start_power_meter()
        self.bind(meter)
        restarted = HeadlessBackend(JsonSettings(self.settings.path))
        try:
            self.assertEqual(restarted.power_meter_profiles.library()["profiles"], [])
            with self.assertRaisesRegex(ValueError, "建立系数方案"):
                restarted.power_meter_profiles.require_ready()
        finally:
            restarted.export_worker.graceful_stop(timeout_s=5.0)
        self.command("create", name="Replacement", softwareGain=1.1, wavelengthNm=977.0)
        self.backend.power_meter_profiles.require_ready()

    def test_delete_rejects_stale_version_wrong_probe_and_missing_profile(self) -> None:
        self.bind(self.option("PROBE-A"))
        profile_id = self.library()["selectedProfileId"]
        self.command("save", profileId=profile_id, expectedVersion=1,
                     name="Updated", softwareGain=1.1, wavelengthNm=977.0)
        before = copy.deepcopy(self.backend.device_profiles)
        for params, message in [
            ({"profileId": profile_id, "expectedVersion": 1}, "刷新"),
            ({"profileId": profile_id, "expectedVersion": 2, "serialNumber": "OTHER"}, "探头已切换"),
            ({"profileId": "missing", "expectedVersion": 2}, "方案不存在"),
        ]:
            with self.subTest(params=params), self.assertRaisesRegex(ValueError, message):
                self.command("delete", **params)
            self.assertEqual(self.backend.device_profiles, before)

    def test_delete_obeys_acquisition_automatic_discovery_and_unsaved_point_locks(self) -> None:
        self.bind(self.option("PROBE-A"))
        profile_id = self.library()["selectedProfileId"]
        before = copy.deepcopy(self.backend.device_profiles)
        for target, attribute, value in [
            (self.backend.runtime, "power_meter_reader", MagicMock()),
            (self.backend.runtime.automatic, "state", AutomaticTestState.WAITING_STABLE),
            (self.backend.runtime, "_automatic_thread", MagicMock()),
            (self.backend, "power_meter_detect_thread", MagicMock()),
            (self.backend.runtime.record_store, "unsaved_records", MagicMock(return_value=[object()])),
        ]:
            with self.subTest(attribute=attribute), patch.object(target, attribute, value):
                with self.assertRaises(ValueError):
                    self.command("delete", profileId=profile_id, expectedVersion=1)
            self.assertEqual(self.backend.device_profiles, before)

    def test_delete_persistence_failure_rolls_back_catalog_configuration_and_disk(self) -> None:
        self.bind(self.option("PROBE-A"))
        profile_id = self.library()["selectedProfileId"]
        before_profiles = copy.deepcopy(self.backend.device_profiles)
        before_config = dict(self.backend.runtime.configuration)
        before_disk = self.settings.path.read_bytes()
        with patch.object(self.settings, "flush", side_effect=OSError("disk full")):
            with self.assertRaisesRegex(OSError, "disk full"):
                self.command("delete", profileId=profile_id, expectedVersion=1)
        self.assertEqual(self.backend.device_profiles, before_profiles)
        self.assertEqual(self.backend.runtime.configuration, before_config)
        self.assertEqual(self.settings.get("runtime/deviceProfiles"), before_profiles)
        self.assertEqual(self.settings.path.read_bytes(), before_disk)

    def test_delete_preserves_frozen_session_measurement_provenance(self) -> None:
        self.bind(self.option("PROBE-A"))
        self.command("create", name="Recorded profile", softwareGain=1.23456789, wavelengthNm=978.25)
        profile_id = self.library()["selectedProfileId"]
        self.backend.runtime.configuration.update({"sn": "RECORDED", "outputDir": self.temp_dir.name})
        self.backend.runtime._ensure_manual_recording_session()
        session = self.backend.runtime.record_store.current_session
        frozen = copy.deepcopy(session.settings)
        self.command("delete", profileId=profile_id, expectedVersion=1)
        self.assertEqual(session.settings, frozen)
        self.assertEqual(session.settings["powerMeterProfileId"], profile_id)
        self.assertEqual(session.settings["softwareGain"], 1.23456789)

    def test_new_probe_gets_a_default_profile_and_snapshot_wire_shape(self) -> None:
        meter = self.option("probe-a")
        snapshot = self.bind(meter)

        config = snapshot["configuration"]
        library = snapshot["devices"]["powerMeter"]["coefficientLibrary"]
        profile = library["profiles"][0]
        revision = profile["versions"][0]
        self.assertEqual(library["serialNumber"], "PROBE-A")
        self.assertEqual(library["driverKind"], "caihuang")
        self.assertEqual(library["selectedProfileId"], profile["id"])
        self.assertEqual(library["selectedVersion"], 1)
        self.assertEqual(revision["name"], "默认方案")
        self.assertEqual(revision["softwareGain"], 1.0)
        self.assertEqual(revision["wavelengthNm"], 976.0)
        self.assertEqual(config["powerMeterProfileId"], profile["id"])
        self.assertEqual(config["powerMeterProfileName"], "默认方案")
        self.assertEqual(config["powerMeterProfileVersion"], 1)
        self.assertFalse(config["powerMeterProfileWavelengthPending"])

        resource_profile = snapshot["devices"]["powerMeter"]["resourceProfiles"][0]
        self.assertEqual(resource_profile["coefficientLibrary"], library)

    def test_identity_includes_driver_and_normalized_probe_serial(self) -> None:
        caihuang = self.option("same-sn", resource="ASRL3::INSTR")
        laserpoint = self.option(
            "same-sn", driver="laserpoint", resource="ASRL7::INSTR"
        )
        self.backend.power_meter_options = [caihuang, laserpoint]
        self.bind(caihuang)
        self.command(
            "create", name="Caihuang calibrated", softwareGain=1.12, wavelengthNm=975.0
        )
        self.bind(laserpoint)

        self.assertEqual(self.backend.runtime.configuration["softwareGain"], 1.0)
        self.assertEqual(self.backend.runtime.configuration["powerMeterProfileName"], "默认方案")
        self.assertIn("caihuang:SAME-SN", self.backend.device_profiles["powerMeters"])
        self.assertIn("laserpoint:SAME-SN", self.backend.device_profiles["powerMeters"])

    def test_probe_without_serial_has_no_library_and_cannot_mutate_profiles(self) -> None:
        self.bind(self.option(""))
        self.assertIsNone(
            self.backend.snapshot()["devices"]["powerMeter"]["coefficientLibrary"]
        )
        with self.assertRaisesRegex(ValueError, "探头 SN"):
            self.command(
                "create", name="No identity", softwareGain=1.0, wavelengthNm=976.0
            )

    def test_create_save_select_history_and_restore_are_append_only(self) -> None:
        self.bind(self.option("PROBE-A"))
        default_id = self.library()["selectedProfileId"]
        created = self.command(
            "create",
            name="976 nm 校准",
            softwareGain=1.12,
            wavelengthNm=976.4,
            notes="first",
            operator="tester-a",
            reason="bench calibration",
        )
        created_library = created["devices"]["powerMeter"]["coefficientLibrary"]
        created_id = created_library["selectedProfileId"]
        self.assertNotEqual(created_id, default_id)
        self.assertEqual(created["configuration"]["softwareGain"], 1.12)

        saved = self.command(
            "save",
            profileId=created_id,
            expectedVersion=1,
            name="976 nm 校准",
            softwareGain=1.18,
            wavelengthNm=976.6,
            notes="second",
            operator="tester-b",
            reason="repeat calibration",
        )
        self.assertEqual(saved["configuration"]["powerMeterProfileVersion"], 2)
        history = self.command("history", profileId=created_id)
        self.assertEqual([item["version"] for item in history["versions"]], [1, 2])
        self.assertEqual(history["versions"][0]["softwareGain"], 1.12)
        self.assertEqual(history["versions"][1]["softwareGain"], 1.18)

        selected = self.command("select", profileId=default_id, expectedVersion=1)
        self.assertEqual(selected["configuration"]["powerMeterProfileName"], "默认方案")
        restored = self.command(
            "restore",
            profileId=created_id,
            expectedVersion=2,
            restoreVersion=1,
            operator="tester-c",
            reason="restore accepted calibration",
        )
        self.assertEqual(restored["configuration"]["powerMeterProfileVersion"], 3)
        self.assertEqual(restored["configuration"]["softwareGain"], 1.12)
        restored_history = self.command("history", profileId=created_id)
        self.assertEqual([item["version"] for item in restored_history["versions"]], [1, 2, 3])
        self.assertEqual(restored_history["versions"][0], history["versions"][0])
        self.assertEqual(restored_history["versions"][2]["reason"], "restore accepted calibration")

    def test_invalid_values_duplicate_names_and_stale_writes_do_not_change_catalog(self) -> None:
        self.bind(self.option("PROBE-A"))
        original = copy.deepcopy(self.backend.device_profiles)
        invalid_values = (
            {"name": "", "softwareGain": 1.0, "wavelengthNm": 976.0},
            {"name": "bad gain", "softwareGain": 0, "wavelengthNm": 976.0},
            {"name": "bad wave", "softwareGain": 1.0, "wavelengthNm": math.nan},
        )
        for values in invalid_values:
            with self.subTest(values=values), self.assertRaises(ValueError):
                self.command("create", **values)
            self.assertEqual(self.backend.device_profiles, original)

        created = self.command(
            "create", name="A", softwareGain=1.1, wavelengthNm=976.0
        )
        profile_id = created["configuration"]["powerMeterProfileId"]
        before_duplicate = copy.deepcopy(self.backend.device_profiles)
        with self.assertRaisesRegex(ValueError, "同名"):
            self.command("create", name="A", softwareGain=1.2, wavelengthNm=977.0)
        self.assertEqual(self.backend.device_profiles, before_duplicate)

        self.command(
            "save",
            profileId=profile_id,
            expectedVersion=1,
            name="A",
            softwareGain=1.2,
            wavelengthNm=977.0,
        )
        before_stale = copy.deepcopy(self.backend.device_profiles)
        with self.assertRaisesRegex(ValueError, "刷新"):
            self.command(
                "save",
                profileId=profile_id,
                expectedVersion=1,
                name="A",
                softwareGain=9.9,
                wavelengthNm=999.0,
            )
        self.assertEqual(self.backend.device_profiles, before_stale)

    def test_persistence_failure_rolls_back_catalog_and_active_configuration(self) -> None:
        self.bind(self.option("PROBE-A"))
        created = self.command(
            "create", name="rollback", softwareGain=1.1, wavelengthNm=976.0
        )
        profile_id = created["configuration"]["powerMeterProfileId"]
        before_profiles = copy.deepcopy(self.backend.device_profiles)
        before_config = dict(self.backend.runtime.configuration)

        with patch.object(
            self.backend, "_persist_configuration", side_effect=OSError("disk full")
        ):
            with self.assertRaisesRegex(OSError, "disk full"):
                self.command(
                    "save",
                    profileId=profile_id,
                    expectedVersion=1,
                    name="rollback",
                    softwareGain=1.25,
                    wavelengthNm=977.0,
                )

        self.assertEqual(self.backend.device_profiles, before_profiles)
        self.assertEqual(self.backend.runtime.configuration, before_config)

    def test_concurrent_persist_waits_for_failed_profile_transaction_to_roll_back(self) -> None:
        self.bind(self.option("PROBE-A"))
        created = self.command(
            "create", name="atomic", softwareGain=1.1, wavelengthNm=976.0
        )
        profile_id = created["configuration"]["powerMeterProfileId"]
        before_profiles = copy.deepcopy(self.backend.device_profiles)
        before_config = dict(self.backend.runtime.configuration)
        before_disk = json.loads(self.settings.path.read_text(encoding="utf-8"))

        profile_flush_entered = threading.Event()
        allow_profile_flush_to_fail = threading.Event()
        concurrent_started = threading.Event()
        concurrent_finished = threading.Event()
        profile_errors: list[Exception] = []
        concurrent_errors: list[Exception] = []
        original_flush = self.settings.flush
        flush_calls = 0

        def controlled_flush() -> None:
            nonlocal flush_calls
            flush_calls += 1
            if flush_calls == 1:
                profile_flush_entered.set()
                if not allow_profile_flush_to_fail.wait(2.0):
                    raise AssertionError("timed out waiting to release the profile flush")
                raise OSError("disk full")
            original_flush()

        def save_profile() -> None:
            try:
                self.command(
                    "save",
                    profileId=profile_id,
                    expectedVersion=1,
                    name="atomic",
                    softwareGain=9.9,
                    wavelengthNm=999.0,
                )
            except Exception as exc:
                profile_errors.append(exc)

        def persist_concurrently() -> None:
            concurrent_started.set()
            try:
                self.backend._persist_configuration()
            except Exception as exc:
                concurrent_errors.append(exc)
            finally:
                concurrent_finished.set()

        with patch.object(self.settings, "flush", side_effect=controlled_flush):
            profile_thread = threading.Thread(target=save_profile)
            profile_thread.start()
            self.assertTrue(profile_flush_entered.wait(2.0))
            concurrent_thread = threading.Thread(target=persist_concurrently)
            concurrent_thread.start()
            self.assertTrue(concurrent_started.wait(2.0))
            self.assertFalse(concurrent_finished.is_set())
            allow_profile_flush_to_fail.set()
            profile_thread.join(2.0)
            concurrent_thread.join(2.0)

        self.assertFalse(profile_thread.is_alive())
        self.assertFalse(concurrent_thread.is_alive())
        self.assertEqual(len(profile_errors), 1)
        self.assertRegex(str(profile_errors[0]), "disk full")
        self.assertEqual(concurrent_errors, [])
        self.assertEqual(self.backend.device_profiles, before_profiles)
        self.assertEqual(self.backend.runtime.configuration, before_config)
        self.assertEqual(self.settings.get("runtime/deviceProfiles"), before_profiles)
        self.assertEqual(self.settings.get("runtime/configuration"), before_config)
        after_disk = json.loads(self.settings.path.read_text(encoding="utf-8"))
        self.assertEqual(after_disk, before_disk)

    def test_ordinary_configuration_cannot_overwrite_profile_provenance(self) -> None:
        meter = self.option("PROBE-A")
        snapshot = self.bind(meter)
        before = dict(snapshot["configuration"])

        with self.assertRaisesRegex(ValueError, "系数方案"):
            self.backend.dispatch(
                "app.configure",
                {"powerMeterResource": meter.resource, "softwareGain": 2.0},
            )
        with self.assertRaisesRegex(ValueError, "系数方案"):
            self.backend.dispatch(
                "app.configure",
                {"powerMeterResource": meter.resource, "powerMeterWavelengthNm": 980.0},
            )

        self.backend.dispatch(
            "app.configure",
            {
                "powerMeterProfileId": "forged",
                "powerMeterProfileName": "forged",
                "powerMeterProfileVersion": 99,
                "powerMeterProfileNotes": "forged",
                "powerMeterProfileWavelengthPending": True,
            },
        )
        after = self.backend.runtime.configuration
        for key in (
            "powerMeterProfileId",
            "powerMeterProfileName",
            "powerMeterProfileVersion",
            "powerMeterProfileNotes",
            "powerMeterProfileWavelengthPending",
            "softwareGain",
            "powerMeterWavelengthNm",
        ):
            self.assertEqual(after[key], before[key])

    def test_switching_probe_restores_target_profile_and_ignores_incoming_gain(self) -> None:
        meter_a = self.option("PROBE-A", resource="ASRL3::INSTR")
        meter_b = self.option("PROBE-B", resource="ASRL7::INSTR")
        self.backend.power_meter_options = [meter_a, meter_b]
        self.bind(meter_a)
        self.command("create", name="A profile", softwareGain=1.31, wavelengthNm=976.1)
        self.bind(meter_b)
        self.command("create", name="B profile", softwareGain=0.82, wavelengthNm=977.2)
        self.bind(meter_a)

        frontend_payload = dict(self.backend.runtime.configuration)
        frontend_payload["powerMeterResource"] = meter_b.resource
        frontend_payload["softwareGain"] = 9.99
        frontend_payload["powerMeterWavelengthNm"] = 999.0
        switched = self.backend.dispatch("app.configure", frontend_payload)

        config = switched["configuration"]
        self.assertEqual(config["powerMeterSerialNumber"], "PROBE-B")
        self.assertEqual(config["powerMeterProfileName"], "B profile")
        self.assertEqual(config["softwareGain"], 0.82)
        self.assertEqual(config["powerMeterWavelengthNm"], 977.2)

    def test_profile_mutations_are_blocked_during_collection_and_with_pending_records(self) -> None:
        self.bind(self.option("PROBE-A"))
        create = {
            "name": "blocked",
            "softwareGain": 1.0,
            "wavelengthNm": 976.0,
        }

        self.backend.runtime.power_meter_reader = MagicMock()
        with self.assertRaisesRegex(ValueError, "停止采集"):
            self.command("create", **create)
        self.backend.runtime.power_meter_reader = None

        self.backend.runtime.automatic.state = AutomaticTestState.WAITING_STABLE
        with self.assertRaisesRegex(ValueError, "自动测试"):
            self.command("create", **create)
        self.backend.runtime.automatic.state = AutomaticTestState.IDLE

        with patch.object(
            self.backend.runtime.record_store, "unsaved_records", return_value=(object(),)
        ):
            with self.assertRaisesRegex(ValueError, "未保存"):
                self.command("create", **create)

        self.backend.power_meter_detect_thread = MagicMock()
        with self.assertRaisesRegex(ValueError, "正在识别"):
            self.command("create", **create)
        with self.assertRaisesRegex(ValueError, "正在识别"):
            self.backend.power_meter_profiles.require_ready()
        self.backend.power_meter_detect_thread = None

    def test_session_keeps_the_profile_revision_present_when_measurement_started(self) -> None:
        self.bind(self.option("PROBE-A"))
        first = self.command(
            "create", name="Frozen", softwareGain=1.21, wavelengthNm=976.3
        )
        profile_id = first["configuration"]["powerMeterProfileId"]
        self.backend.runtime.configuration.update(
            {"sn": "SESSION-SN", "outputDir": self.temp_dir.name}
        )
        self.backend.runtime._ensure_manual_recording_session()
        session = self.backend.runtime.record_store.current_session
        self.assertIsNotNone(session)
        frozen_settings = dict(session.settings)

        self.command(
            "save",
            profileId=profile_id,
            expectedVersion=1,
            name="Frozen",
            softwareGain=1.35,
            wavelengthNm=977.0,
        )

        self.assertEqual(frozen_settings["powerMeterProfileId"], profile_id)
        self.assertEqual(frozen_settings["powerMeterProfileVersion"], 1)
        self.assertEqual(frozen_settings["softwareGain"], 1.21)
        self.assertEqual(frozen_settings["powerMeterWavelengthNm"], 976.3)
        self.assertEqual(dict(session.settings), frozen_settings)

    def test_manual_save_keeps_measurement_profile_but_accepts_latest_destination_metadata(self) -> None:
        self.bind(self.option("PROBE-A", driver="coherent"))
        created = self.command(
            "create",
            name="Frozen manual profile",
            softwareGain=1.21,
            wavelengthNm=976.3,
            notes="fixture A",
        )
        profile_id = created["configuration"]["powerMeterProfileId"]
        initial_output = Path(self.temp_dir.name) / "initial"
        saved_output = Path(self.temp_dir.name) / "saved"
        self.backend.runtime.configuration.update(
            {
                "sn": "MEASURED-SN",
                "station": "初测",
                "outputDir": str(initial_output),
            }
        )
        self.backend.runtime._ensure_manual_recording_session()
        original_session = self.backend.runtime.record_store.current_session
        self.assertIsNotNone(original_session)
        self.backend.runtime.record_store.queue(
            ExcelTestRecord(
                current_a=1.0,
                voltage_v=2.0,
                power_w=1.0,
                efficiency=0.5,
                peak_wavelength_nm=976.2,
                centroid_nm=976.2,
                fwhm_nm=0.4,
                pib=0.98,
                wavelength=[976.0, 976.2],
                intensity=[10.0, 20.0],
                test_station="初测",
            )
        )

        # Simulate unrelated configuration changes after the accepted point.
        # Protected measurement fields must still come from the session start.
        self.backend.runtime.configuration.update(
            {
                "sn": "SAVED-SN",
                "station": "复测",
                "outputDir": str(saved_output),
                "powerMeterDriverKind": "laserpoint",
                "powerMeterSerialNumber": "OTHER-PROBE",
                "powerMeterProfileId": "forged",
                "powerMeterProfileName": "later profile",
                "powerMeterProfileVersion": 99,
                "powerMeterProfileNotes": "later fixture",
                "powerMeterProfileWavelengthPending": True,
                "softwareGain": 9.9,
                "powerMeterWavelengthNm": 999.0,
            }
        )
        exported_path = saved_output / "result.xlsx"
        with patch.object(
            self.backend.runtime, "_save_record_store", return_value=exported_path
        ):
            result = self.backend.runtime.save_manual_records()

        saved_session = self.backend.runtime.record_store.current_session
        self.assertIsNotNone(saved_session)
        self.assertEqual(result, {"outputPath": str(exported_path)})
        self.assertEqual(saved_session.sn, "SAVED-SN")
        self.assertEqual(saved_session.station, "复测")
        self.assertTrue(saved_output in saved_session.workbook_path.parents)
        self.assertEqual(saved_session.settings["sn"], "SAVED-SN")
        self.assertEqual(saved_session.settings["station"], "复测")
        self.assertEqual(saved_session.settings["outputDir"], str(saved_output))
        self.assertEqual(saved_session.settings["powerMeterDriverKind"], "coherent")
        self.assertEqual(saved_session.settings["powerMeterSerialNumber"], "PROBE-A")
        self.assertEqual(saved_session.settings["powerMeterProfileId"], profile_id)
        self.assertEqual(
            saved_session.settings["powerMeterProfileName"], "Frozen manual profile"
        )
        self.assertEqual(saved_session.settings["powerMeterProfileVersion"], 1)
        self.assertEqual(saved_session.settings["powerMeterProfileNotes"], "fixture A")
        self.assertFalse(saved_session.settings["powerMeterProfileWavelengthPending"])
        self.assertEqual(saved_session.settings["softwareGain"], 1.21)
        self.assertEqual(saved_session.settings["powerMeterWavelengthNm"], 976.3)
        self.assertEqual(
            self.backend.runtime.record_store.snapshot()[0].test_station, "复测"
        )

    def test_export_writes_selected_profile_provenance_to_chinese_session_rows(self) -> None:
        self.bind(self.option("PROBE-A", driver="coherent"))
        self.command(
            "create",
            name="976 nm 工装校准",
            softwareGain=1.21,
            wavelengthNm=976.3,
            notes="25 C, fixture A",
        )
        self.backend.runtime.configuration.update(
            {"sn": "EXPORT-SN", "outputDir": self.temp_dir.name}
        )
        self.backend.runtime._ensure_manual_recording_session()
        session = self.backend.runtime.record_store.current_session
        self.assertIsNotNone(session)
        record = ExcelTestRecord(
            current_a=1.0,
            voltage_v=2.0,
            power_w=1.0,
            efficiency=0.5,
            peak_wavelength_nm=976.2,
            centroid_nm=976.2,
            fwhm_nm=0.4,
            pib=0.98,
            wavelength=[976.0, 976.2],
            intensity=[10.0, 20.0],
        )
        target = save_test_records(session.workbook_path, [record], session=session)

        workbook = load_workbook(target, read_only=True, data_only=True)
        try:
            sheet = workbook["会话信息"]
            rows = {
                str(row[0].value): row[1].value
                for row in sheet.iter_rows(min_col=1, max_col=2)
                if row[0].value is not None
            }
        finally:
            workbook.close()
        self.assertEqual(rows["功率计探头 SN"], "PROBE-A")
        self.assertEqual(rows["功率计驱动"], "coherent")
        self.assertEqual(rows["系数方案"], "976 nm 工装校准")
        self.assertEqual(rows["系数方案版本"], "1")
        self.assertEqual(rows["功率系数"], "1.21")
        self.assertEqual(rows["功率计校准波长 (nm)"], "976.3")
        self.assertEqual(rows["方案条件备注"], "25 C, fixture A")


class PowerMeterProfileMigrationTests(unittest.TestCase):
    def test_legacy_profiles_migrate_once_and_only_current_probe_inherits_wavelength(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            settings = JsonSettings(Path(temp_dir) / "settings.json")
            settings.update(
                {
                    "runtime/configuration": {
                        "powerMeterResource": "ASRL3::INSTR",
                        "powerMeterDriverKind": "coherent",
                        "powerMeterSerialNumber": "HEAD-A",
                        "softwareGain": 1.25,
                        "powerMeterWavelengthNm": 980.0,
                    },
                    "runtime/deviceProfiles": {
                        "powerMeters": {
                            "coherent:HEAD-A": {
                                "driverKind": "coherent",
                                "serialNumber": "HEAD-A",
                                "softwareGain": 1.25,
                            },
                            "coherent:HEAD-B": {
                                "driverKind": "coherent",
                                "serialNumber": "HEAD-B",
                                "softwareGain": 0.84,
                            },
                        }
                    },
                }
            )
            backend = HeadlessBackend(settings)
            try:
                current = backend.device_profiles["powerMeters"]["coherent:HEAD-A"]
                other = backend.device_profiles["powerMeters"]["coherent:HEAD-B"]
                current_revision = current["profiles"][0]["versions"][0]
                other_revision = other["profiles"][0]["versions"][0]
                self.assertEqual(current_revision["name"], "原有系数")
                self.assertEqual(current_revision["softwareGain"], 1.25)
                self.assertEqual(current_revision["wavelengthNm"], 980.0)
                self.assertEqual(other_revision["name"], "原有系数")
                self.assertEqual(other_revision["softwareGain"], 0.84)
                self.assertIsNone(other_revision["wavelengthNm"])
                first_persisted = copy.deepcopy(settings.get("runtime/deviceProfiles"))
            finally:
                backend.export_worker.graceful_stop(timeout_s=5.0)

            restarted = HeadlessBackend(JsonSettings(settings.path))
            try:
                self.assertEqual(restarted.device_profiles, first_persisted)
                self.assertEqual(
                    restarted.runtime.configuration["powerMeterProfileName"], "原有系数"
                )
                self.assertEqual(restarted.runtime.configuration["softwareGain"], 1.25)
                self.assertEqual(
                    restarted.runtime.configuration["powerMeterWavelengthNm"], 980.0
                )
            finally:
                restarted.export_worker.graceful_stop(timeout_s=5.0)

    def test_migrated_probe_with_unknown_wavelength_blocks_collection_until_saved(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            settings = JsonSettings(Path(temp_dir) / "settings.json")
            settings.update(
                {
                    "runtime/configuration": {
                        "powerMeterResource": "ASRL9::INSTR",
                        "powerMeterDriverKind": "coherent",
                        "powerMeterSerialNumber": "HEAD-B",
                    },
                    "runtime/deviceProfiles": {
                        "powerMeters": {
                            "coherent:HEAD-A": {
                                "driverKind": "coherent",
                                "serialNumber": "HEAD-A",
                                "softwareGain": 1.25,
                            },
                            "coherent:HEAD-B": {
                                "driverKind": "coherent",
                                "serialNumber": "HEAD-B",
                                "softwareGain": 0.84,
                            },
                        }
                    },
                }
            )
            backend = HeadlessBackend(settings)
            try:
                # The default 976 nm value is not proof that the old user chose
                # it. With no persisted wavelength, even the current legacy
                # probe must remain pending until the operator saves a revision.
                self.assertTrue(
                    backend.runtime.configuration["powerMeterProfileWavelengthPending"]
                )
                with self.assertRaisesRegex(ValueError, "校准波长待补充"):
                    backend.power_meter_profiles.require_ready()
            finally:
                backend.export_worker.graceful_stop(timeout_s=5.0)


if __name__ == "__main__":
    unittest.main()
