"""Hardware-free regressions for the safety and performance audit."""

import copy
import json
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from combined_test.excel_export import append_test_record, save_test_records
from combined_test.automation import AutomaticTestSettings, AutomaticTestState
from combined_test.record_store import SessionRecordStore
from tauri_bridge.headless_backend import HeadlessBackend
from tauri_bridge.service import BridgeService
from tauri_bridge.settings import JsonSettings
from tests.test_station_runtime import _FakeController
from tests import test_excel_export


class SafetyRegressions(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.backend = HeadlessBackend(JsonSettings(Path(self.temp.name) / "settings.json"))
        self.addCleanup(self.backend.export_worker.graceful_stop, 2)
        self.station = self.backend.runtime
        self.controller = _FakeController()
        self.station.power_supply_controller = self.controller
        self.station.power_supply_kind = "tdk"
        self.station.configuration["powerSupplyKind"] = "tdk"

    def test_queued_command_is_rejected_after_stop_and_new_command_is_allowed(self):
        service = BridgeService(self.backend)
        queued = self.backend.capture_command_generation()
        self.station.end_automatic()
        request = json.dumps({"v": 2, "id": "old", "method": "powerSupply.setCurrent", "params": {"currentA": 5}})
        response = service.handle_line(request, command_generation=queued)
        self.assertIn("error", response)
        self.assertEqual(self.controller.current_a, 0)
        self.assertTrue(self.station.shutdown_confirmed)
        self.backend.dispatch("powerSupply.setCurrent", {"currentA": 5})
        self.assertEqual(self.controller.current_a, 5)
        self.assertFalse(self.station.shutdown_confirmed)

    def test_commands_captured_during_stop_never_become_valid(self):
        with self.station.command_safety.stopping():
            queued = self.backend.capture_command_generation()
            with self.assertRaises(RuntimeError):
                self.station.set_current({"currentA": 1})
        with self.assertRaises(RuntimeError):
            self.backend.dispatch("powerSupply.setCurrent", {"currentA": 1}, command_generation=queued)

    def test_stop_is_rechecked_after_waiting_for_io(self):
        def reserve(**kwargs):
            # A stop can invalidate a command while it is reserving its device
            # timing slot. The subsequent I/O check must reject that command.
            with self.station.command_safety.stopping():
                pass
        with patch.object(self.station, "_reserve_power_supply_command", reserve):
            with self.assertRaises(RuntimeError):
                self.station.set_current({"currentA": 5})
        self.assertEqual(self.controller.writes, [])

    def test_disconnect_waits_for_ramp_lock_without_holding_station_or_io_lock(self):
        self.station.power_supply_kind = "ch341"
        self.station.active_current_a = 1
        entered = threading.Event()
        finished = threading.Event()
        errors = []
        original_lock = self.station._safe_ramp_lock
        class ObservedRampLock:
            def __enter__(self):
                if threading.current_thread().name == "audit-disconnect":
                    entered.set()
                return original_lock.__enter__()
            def __exit__(self, *args):
                return original_lock.__exit__(*args)
        self.station._safe_ramp_lock = ObservedRampLock()
        def disconnect():
            try:
                self.station.disconnect_power_supply()
            except Exception as exc:
                errors.append(exc)
            finally:
                finished.set()

        with self.station._safe_ramp_lock:
            worker = threading.Thread(target=disconnect, name="audit-disconnect", daemon=True)
            worker.start()
            self.assertTrue(entered.wait(1))
            acquired = self.station._lock.acquire(timeout=0.2)
            self.assertTrue(acquired, "disconnect must not hold state while waiting for ramp")
            if acquired:
                self.station._lock.release()
            acquired = self.station._supply_io_lock.acquire(timeout=0.2)
            self.assertTrue(acquired)
            if acquired:
                self.station._supply_io_lock.release()
        self.assertTrue(finished.wait(2))
        self.assertFalse(errors)
        self.assertFalse(self.controller.is_connected)

    def test_stop_during_automatic_start_cannot_clear_stop_or_launch_worker(self):
        def start_meter(**kwargs):
            self.station.request_emergency_stop()
            self.station._emergency_ramp_thread.join(2)
        with patch.object(self.station, "_require_automatic_ready", return_value=AutomaticTestSettings(use_spectrometer=False)), \
             patch.object(self.station, "start_power_meter", side_effect=start_meter):
            with self.assertRaises(RuntimeError):
                self.station.start_automatic({"sn": "SN", "outputDir": self.temp.name})
        self.assertIsNone(self.station._automatic_thread)
        self.assertTrue(self.station._automatic_stop.is_set())
        self.assertEqual(self.station.automatic.state, AutomaticTestState.COMPLETED)
        self.assertTrue(all(value == 0 for value in self.controller.writes))

    def test_invalid_pd_configuration_does_not_change_settings(self):
        before = copy.deepcopy(self.station.pd_state["settings"])
        for params in ({"sampleRateHz": 0}, {"save": "false"}, {"blockSize": 1.5}, {"scale": float("inf")}):
            with self.subTest(params=params), self.assertRaises(ValueError):
                self.backend.dispatch("pd.configure", params)
            self.assertEqual(self.station.pd_state["settings"], before)

    def test_invalid_configuration_is_atomic(self):
        for values in ({"sn": "bad", "currentPlanMode": "invalid"}, {"sn": "bad", "stableWindowS": 0},
                       {"sn": "bad", "autoIntegration": "false"}, {"targetCurrentA": float("nan")},
                       {"customCurrentPoints": [{"currentA": 1, "dwellS": -1}]}):
            before = copy.deepcopy(self.station.configuration)
            with self.subTest(values=values), self.assertRaises(ValueError):
                self.backend.dispatch("app.configure", values)
            self.assertEqual(before, self.station.configuration)

    def test_invalid_hardware_parameters_never_reach_driver(self):
        for value in ("false", "true", 0, 1, None):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.station.set_output({"enabled": value})
        for value in (True, "5", float("nan"), float("inf"), -1):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.station.set_current({"currentA": value})
        self.assertEqual(self.controller.writes, [])

    def test_missing_parameters_are_reported_as_invalid_params(self):
        service = BridgeService(self.backend)
        for method in ("powerSupply.setCurrent", "powerSupply.setVoltage", "powerSupply.setOutput"):
            reply = service.handle_line(json.dumps({"v": 2, "id": "missing", "method": method, "params": {}}))
            self.assertEqual(reply["error"]["code"], "invalid_params")

    def test_failed_output_write_clears_previous_safe_confirmation(self):
        self.station.safe_ramp_down()
        with patch.object(self.controller, "set_output_current", side_effect=RuntimeError("no response")):
            with self.assertRaises(RuntimeError):
                self.station.set_current({"currentA": 1})
        self.assertFalse(self.station.shutdown_confirmed)

    def test_settings_are_batched_and_runtime_commands_do_not_flush(self):
        with patch.object(self.backend.settings, "flush") as flush:
            self.backend.dispatch("app.configure", {"sn": "one"})
            self.assertEqual(flush.call_count, 1)
            self.backend.dispatch("app.configure", {"sn": "one"})
            self.backend.dispatch("charts.reset", {})
            self.backend.dispatch("powerSupply.read", {"value": "outputCurrent"})
            self.assertEqual(flush.call_count, 1)

    def test_hidden_and_unchanged_pd_history_is_not_traversed(self):
        class UnreadableHistory(list):
            def __iter__(self):
                raise AssertionError("unneeded PD history copy")

        self.station.pd_state["points"] = UnreadableHistory([{}])
        self.backend.snapshot({"view": "automatic"})
        self.backend.snapshot({"view": "pd", "since": {"pd": self.station.series_revisions["pd"]}})


class RecordSafetyRegressions(unittest.TestCase):
    def test_session_paths_are_reserved_before_export_and_reused_for_retry(self):
        with tempfile.TemporaryDirectory() as root:
            stores = [SessionRecordStore() for _ in range(3)]
            when = datetime(2026, 9, 7, 12, 34)
            paths = [store.begin_session(Path(root), "SN", when, test_station="ST") for store in stores]
            self.assertEqual([p.stem for p in paths], ["2026_09_07_12_34", "2026_09_07_12_34_02", "2026_09_07_12_34_03"])
            self.assertEqual(stores[1].retarget_workbook(Path(root)), paths[1])
            self.assertEqual(stores[1].retarget_manual_session(Path(root), "SN", "ST"), paths[1])
            self.assertEqual(stores[1].begin_session(Path(root), "SN", when, reset=False), paths[1])

    def test_concurrent_sessions_and_existing_workbooks_are_not_overwritten(self):
        with tempfile.TemporaryDirectory() as root:
            existing = Path(root) / "SN" / "ST" / "2026_09_07_12_34.xlsx"
            existing.parent.mkdir(parents=True)
            existing.write_bytes(b"existing record")
            def reserve(_):
                return SessionRecordStore().begin_session(Path(root), "SN", datetime(2026, 9, 7, 12, 34), test_station="ST")
            with ThreadPoolExecutor(max_workers=4) as pool:
                paths = list(pool.map(reserve, range(8)))
            self.assertEqual(len(set(paths)), 8)
            self.assertNotIn(existing, paths)
            self.assertEqual(existing.read_bytes(), b"existing record")

    def test_user_text_is_never_exported_as_formula(self):
        with tempfile.TemporaryDirectory() as root:
            record = test_excel_export.ExcelExportTests()._record(2, test_station="=1+1")
            store = SessionRecordStore()
            path = store.begin_session(Path(root), "=1+1", datetime.now(), test_station="=1+1", settings={"testTemperatureC": "=1+1"})
            save_test_records(path, [record], session=store.current_session)
            workbook = load_workbook(path, data_only=False)
            try:
                texts = [cell for sheet in workbook for row in sheet for cell in row if cell.value == "=1+1"]
                self.assertGreaterEqual(len(texts), 3)
                self.assertTrue(all(cell.data_type == "s" for cell in texts))
            finally:
                workbook.close()
            append_path = Path(root) / "append.xlsx"
            append_test_record(append_path, record)
            workbook = load_workbook(append_path, data_only=False)
            try:
                self.assertEqual(workbook.active["C1"].data_type, "s")
            finally:
                workbook.close()
