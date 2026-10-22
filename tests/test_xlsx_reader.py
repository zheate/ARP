"""Differential checks against the pinned, unmodified openpyxl reader."""
from __future__ import annotations

import hashlib
import random
import re
import struct
import tempfile
import unittest
import warnings
from datetime import date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree.ElementTree import Element, SubElement

import openpyxl
import xlsxwriter
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from openpyxl.worksheet import _reader

from combined_test import data_browser as browser


def original_rows(sheet, *, min_row, max_col):
    return sheet.iter_rows(min_row=min_row, max_col=max_col, values_only=True)


def exact(value):
    """Equality including Python types, negative zero and every binary64 bit."""
    if isinstance(value, float):
        return (float, struct.pack(">d", value))
    if isinstance(value, dict):
        return (dict, tuple((key, exact(item)) for key, item in value.items()))
    if isinstance(value, (tuple, list)):
        return (type(value), tuple(exact(item) for item in value))
    return (type(value), value)


def outcome(operation):
    with warnings.catch_warnings(record=True) as emitted:
        warnings.simplefilter("always")
        # GC can report an earlier openpyxl failed-load handle during either
        # run. Parser/date warnings remain part of the differential assertion.
        warnings.simplefilter("ignore", ResourceWarning)
        try:
            result = ("ok", exact(operation()))
        except Exception as error:
            result = ("error", type(error), str(error))
        return result, [(item.category, str(item.message)) for item in emitted]


class NumericCoordinateRegression(unittest.TestCase):
    def test_data_browser_reuses_numeric_column_coordinates(self):
        # max_col=2 must still validate all 32 columns. General coordinate
        # parsing per numeric cell is the measured bottleneck, not the output.
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "spectrum.xlsx"
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("测试数据")
            sheet.append(["电流(A)", "功率(W)"])
            for row in range(256):
                sheet.append([float(row + col) for col in range(32)])
            workbook.save(path)
            workbook.close()
            with patch.object(_reader, "coordinate_to_tuple", wraps=_reader.coordinate_to_tuple) as convert:
                detail = browser._read_workbook(path, "test", 0)
            self.assertEqual(len(detail["points"]), 256)
            self.assertLess(convert.call_count, 100, "numeric columns were reparsed for every cell")


class XlsxReaderCompatibility(unittest.TestCase):
    def setUp(self):
        from combined_test.xlsx_reader import iter_browser_rows
        self.read_rows = iter_browser_rows
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        browser.clear_data_browser_cache()
        self.addCleanup(browser.clear_data_browser_cache)
        self.base = self.root / "base.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "测试数据"
        sheet.append(["SN", "LEGACY", "测试站别", "封盖前"])
        sheet.append(["电流(A)", "功率(W)", "电光效率", "2 A", "强度"])
        sheet.append([2, 0.125, 0.5, 975.1, -0.0])
        sheet.append([4, 1e-300, 0.25, 976.2, 12345.75])
        sheet.append([None, None, None, 977.3, 4.5])
        workbook.save(self.base)
        workbook.close()

    def mutated(self, name, transform, member="xl/worksheets/sheet1.xml"):
        path = self.root / (name + ".xlsx")
        with ZipFile(self.base) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == member:
                    changed = transform(data.decode()).encode()
                    self.assertNotEqual(changed, data, "fixture mutation did not apply")
                    data = changed
                target.writestr(item, data)
        return path

    def assert_same(self, path):
        before = hashlib.sha256(path.read_bytes()).digest()
        for minimum, maximum in ((1, 2), (2, 5), (4, 40)):
            def read(iterator):
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    # Keep the yielded prefix too: error position is observable.
                    rows = []
                    result = outcome(lambda: rows.extend(iterator(workbook.worksheets[0], min_row=minimum, max_col=maximum)))
                    return exact(rows), result
                finally:
                    workbook.close()
            self.assertEqual(outcome(lambda: read(original_rows)), outcome(lambda: read(self.read_rows)))
        # Compare the complete product result too: header/layout recognition,
        # summaries, warnings, scalar points, percentage conversion and spectra.
        for full in (False, True):
            with patch.object(browser, "iter_browser_rows", original_rows):
                expected = outcome(lambda: browser._read_workbook(path, "test", 0, full))
            self.assertEqual(outcome(lambda: browser._read_workbook(path, "test", 0, full)), expected)
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

    def test_native_legacy_sparse_and_hidden_metadata(self):
        self.assert_same(self.base)
        workbook = load_workbook(self.base)
        sheet = workbook.active
        sheet.cell(12, 26, 1.2345678901234567)
        sheet.cell(31, 4, 976.1234567890123)
        sheet.merge_cells("J8:L9")
        sheet.row_dimensions[4].hidden = True
        sheet.column_dimensions["B"].hidden = True
        info = workbook.create_sheet("会话信息")
        info.append(["SN", "中文 SN"])
        info.sheet_state = "hidden"
        workbook.create_sheet("测量尝试").append([1, 2, 3])
        path = self.root / "hidden.xlsx"
        workbook.save(path)
        workbook.close()
        self.assert_same(path)

    def test_numeric_lexemes_and_cached_formulas(self):
        values = ["-0.0", "0", "-0", "1", "9007199254740993", "1.2345678901234567",
                  "2.2250738585072014e-308", "4.9406564584124654e-324", "1.7976931348623157e308",
                  "1e309", "NaN", "nan", "INF", "-inf", " 1.25 ", "+2", "1_000", "1E-20",
                  "invalid-number", "1e", "", " ", "0x10"]
        for index, value in enumerate(values):
            with self.subTest(value=value):
                self.assert_same(self.mutated(str(index), lambda xml, v=value: xml.replace('<v>0.125</v>', f'<v>{v}</v>')))
        for formula in ("SUM(A1:A2)", "A3*2", "", '<f t="shared" si="0" ref="B3:B4">A3*2</f>'):
            tag = formula if formula.startswith("<f") else f"<f>{formula}</f>"
            with self.subTest(formula=formula):
                self.assert_same(self.mutated("formula", lambda xml, f=tag: xml.replace('<v>0.125</v>', f + '<v>0.125</v>')))

    def test_all_cell_types_styles_and_date_epochs(self):
        for epoch in (None, CALENDAR_MAC_1904):
            workbook = load_workbook(self.base)
            if epoch is not None:
                workbook.epoch = epoch
            sheet = workbook.active
            sheet.append([date(2026, 10, 15), datetime(2026, 10, 15, 12, 34, 56), time(1, 2, 3), timedelta(days=1, seconds=2)])
            sheet.append([True, False, "中文 & < > 😀", "#DIV/0!", "=A3*2"])
            sheet.cell(8, 2, 976.1234567890123).number_format = "0.0000000000000000"
            sheet.cell(9, 2, 40000).number_format = "yyyy-mm-dd"
            sheet.cell(10, 2, 1e100).number_format = "yyyy-mm-dd"
            path = self.root / ("epoch" + str(epoch is not None) + ".xlsx")
            workbook.save(path)
            workbook.close()
            self.assert_same(path)
        for kind, value in (("b", "1"), ("b", "0"), ("b", "invalid"), ("e", "#N/A"), ("str", "cached"),
                            ("d", "2026-10-15T12:34:56"), ("d", "invalid-date"), ("s", "9999"),
                            ("n", "invalid-number"), ("unknown", "text")):
            with self.subTest(kind=kind, value=value):
                self.assert_same(self.mutated("typed", lambda xml, k=kind, v=value: xml.replace('r="B3" t="n"><v>0.125', f'r="B3" t="{k}"><v>{v}')))
        for style in ("0", "00", "+0", "-1", "1", "99", "invalid"):
            with self.subTest(style=style):
                self.assert_same(self.mutated("style", lambda xml, s=style: xml.replace('r="B3"', f'r="B3" s="{s}"')))
        # Default style 0 can itself represent a date. It must disable the shortcut.
        self.assert_same(self.mutated("default-date", lambda xml: xml.replace('<xf numFmtId="0"', '<xf numFmtId="14"'), "xl/styles.xml"))

    def test_coordinates_rows_dimensions_and_unused_invalid_cells(self):
        replacements = [
            ('r="B3"', 'r="b3"'), ('r="B3"', ''), ('r="B3"', 'r="B03"'),
            ('r="B3"', 'r="B4"'), ('r="B3"', 'r="ZZZ3"'), ('r="B3"', 'r="ZZZZ3"'),
            ('r="B3"', 'r="$B$3"'), ('r="B3"', 'r="B-3"'), ('r="B3"', 'r="Ｂ3"'),
            ('r="B3"', 'r="B٣"'), ('r="B3"', 'r="B0"'), ('r="B3"', 'r="B3.0"'),
            ('<row r="3">', '<row>'), ('<row r="3">', '<row r="3.0">'),
            ('<row r="3">', '<row r="3.5">'), ('<row r="3">', '<row r="0">'),
            ('<row r="3">', '<row r="-3">'), ('<row r="3">', '<row r="1">'),
            ('<row r="3">', '<row r="12">'), ('<row r="4">', '<row r="3">'),
            ('ref="A1:E5"', 'ref="A1:B2"'),
            ('ref="A1:E5"', 'ref="A1:Z40"'), ('ref="A1:E5"', 'ref="A1"'),
            ('ref="A1:E5"', 'ref="bad"'),
            # Unused numeric cells must still raise, including in the first row
            # beyond the declared dimension (openpyxl parses before stopping).
            ('</row><row r="4">', '<c r="Z3" t="n"><v>invalid-number</v></c></row><row r="4">'),
            ('</row><row r="4">', '<c r="Z3" t="s"><v>9999</v></c></row><row r="4">'),
            ('</sheetData>', '<row r="6"><c r="Z6" t="n"><v>invalid-number</v></c></row></sheetData>'),
            ('</sheetData>', '<row r="6"><c r="Z6" t="n"><v>1</v></c></row></sheetData>'),
        ]
        for index, (old, new) in enumerate(replacements):
            with self.subTest(old=old, new=new):
                with ZipFile(self.base) as source:
                    self.assertIn(old, source.read("xl/worksheets/sheet1.xml").decode())
                self.assert_same(self.mutated(str(index), lambda xml, a=old, b=new: xml.replace(a, b)))
        self.assert_same(self.mutated("no-dimension", lambda xml: re.sub(r"<dimension [^>]*/>", "", xml)))

    def test_shared_strings_inline_strings_and_writer_modes(self):
        for constant_memory in (False, True):
            path = self.root / f"writer-{constant_memory}.xlsx"
            workbook = xlsxwriter.Workbook(path, {"constant_memory": constant_memory})
            sheet = workbook.add_worksheet("测试数据")
            sheet.write_row(0, 0, ["电流(A)", "功率(W)", "2 A", "强度"])
            sheet.write_row(1, 0, [2, 0.12345678901234567, 976.0123456789012, -0.0])
            sheet.write_formula(2, 0, "=1+1", None, 2)
            sheet.write_formula(2, 1, "=1/8", None, 0.125)
            sheet.write_formula(2, 2, '=IF(1,"cached","")', None, "cached")
            sheet.write_formula(2, 3, "=TRUE()", None, True)
            sheet.write_rich_string(3, 0, "中文", workbook.add_format({"bold": True}), "字符")
            sheet.write_string(3, 1, "_x000D_ & < >")
            workbook.close()
            self.assert_same(path)
        path = self.root / "write-only.xlsx"
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet()
        sheet.append(["电流(A)", "功率(W)"])
        sheet.append([2, 1.2345678901234567])
        workbook.save(path)
        workbook.close()
        self.assert_same(path)

    def test_product_writer_numeric_formats_and_complete_spectra(self):
        from combined_test.excel_export import ExcelTestRecord, save_test_records
        path = self.root / "native-export.xlsx"
        records = [ExcelTestRecord(current_a=current, voltage_v=12.345678901234567,
                                  power_w=1.2345678901234567, efficiency=0.12345678901234567,
                                  peak_wavelength_nm=976.1234567890123, centroid_nm=975.9123456789012,
                                  fwhm_nm=0.12345678901234567, pib=0.9876543210987654,
                                  wavelength=[970 + i * .001 for i in range(128)],
                                  intensity=[-0.0] + [(i + .125) * current for i in range(1, 128)],
                                  test_station="封盖前") for current in (2, 4, 6)]
        save_test_records(path, records)
        self.assert_same(path)
        detail = browser._read_workbook(path, "test", 0, True)
        self.assertEqual([len(item["intensity"]) for item in detail["spectra"]], [128] * 3)

    def test_seeded_cell_combinations_match_original_parser(self):
        from combined_test import xlsx_reader
        if not xlsx_reader._PARSER_AVAILABLE:
            self.skipTest("private parser unavailable; public iterator fallback is covered")
        randomizer = random.Random(20261015)
        for index in range(2000):
            row = randomizer.choice([-1, 0, 1, 3, 10, 100, 1048576])
            coordinate = randomizer.choice([None, f"A{row}", f"Z{row}", f"AA{row}", f"ZZZ{row}",
                                            f"a{row}", f"A0{row}", "B3", "B٣", "A-1", "ZZZZ1", "A1.0", ""])
            kind = randomizer.choice(["n", "n", "n", "n", "s", "b", "str", "e", "d", "inlineStr"])
            style = randomizer.choice([None, "0", "0", "1", "-1", "00", "bad"])
            value = randomizer.choice([None, "", "2", "-0.0", "1.2345678901234567", "1e309", "1e-320",
                                       "9007199254740993", "NaN", "invalid", "2026-10-15"])
            cell = Element("c", {"t": kind})
            if coordinate is not None:
                cell.set("r", coordinate)
            if style is not None:
                cell.set("s", style)
            if value is not None:
                SubElement(cell, _reader.VALUE_TAG).text = value
            parsers = [_reader.WorkSheetParser, xlsx_reader._NumericColumnParser]
            results = []
            for parser_type in parsers:
                parser = parser_type(BytesIO(), ["shared"], data_only=True, date_formats={1})
                parser.row_counter = row
                parser.col_counter = 2
                results.append(outcome(lambda: parser.parse_cell(cell)))
            with self.subTest(index=index, row=row, coordinate=coordinate, kind=kind, style=style, value=value):
                self.assertEqual(results[0], results[1])

    def test_damaged_archive_and_xml_errors(self):
        self.assert_same(self.mutated("xml", lambda xml: xml.replace("</sheetData>", "</broken>")))
        path = self.root / "truncated.xlsx"
        path.write_bytes(self.base.read_bytes()[:-100])
        self.assert_same(path)
        path.write_bytes(b"not a workbook")
        self.assert_same(path)

    def test_index_error_records_filters_refresh_and_repair_match(self):
        invalid = self.mutated("invalid-unused", lambda xml: xml.replace(
            '</row><row r="4">', '<c r="Z3" t="n"><v>invalid-number</v></c></row><row r="4">'))
        with self.assertRaisesRegex(ValueError, "could not convert string to float: 'invalid-number'"):
            browser._read_workbook(invalid, "test", 0)
        (self.root / "broken.xlsx").write_bytes(b"not a workbook")
        (self.root / "~$ignored.xlsx").write_bytes(b"ignored")
        def listings(iterator):
            browser.clear_data_browser_cache()
            with patch.object(browser, "iter_browser_rows", iterator):
                results = []
                for params in ({}, {"refresh": True}, {"status": "read_error"}, {"sn": "legacy"},
                               {"station": "封盖前", "pageSize": 1, "page": 2}):
                    result = browser.list_data_records({"rootPath": str(self.root), **params})
                    result.pop("indexedAtMs")
                    results.append(result)
                return results
        expected = listings(original_rows)
        self.assertEqual(exact(listings(self.read_rows)), exact(expected))
        self.assertEqual(expected[0]["errorCount"], 2)
        # A repaired version must be re-read rather than retaining its error.
        invalid.write_bytes(self.base.read_bytes())
        result = browser.list_data_records({"rootPath": str(self.root), "refresh": True})
        self.assertEqual(result["errorCount"], 1)

    def test_unsupported_versions_and_custom_sheets_use_original_iterator(self):
        from combined_test import xlsx_reader
        with patch.object(openpyxl, "__version__", "future-version"):
            self.assert_same(self.base)
        with patch.object(xlsx_reader, "_PARSER_AVAILABLE", False):
            self.assert_same(self.base)
        class CustomSheet:
            def iter_rows(self, **kwargs):
                self.kwargs = kwargs
                return iter([(1, -0.0)])
        sheet = CustomSheet()
        self.assertEqual(exact(list(self.read_rows(sheet, min_row=2, max_col=2))), exact([(1, -0.0)]))
        self.assertEqual(sheet.kwargs, {"min_row": 2, "max_col": 2, "values_only": True})


if __name__ == "__main__":
    unittest.main()
