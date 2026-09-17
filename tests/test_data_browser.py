from __future__ import annotations

import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from combined_test import data_browser as browser
from tauri_bridge.headless_backend import HeadlessBackend
from tauri_bridge.service import BridgeService, BACKGROUND_FILE_TASK_METHODS
from tauri_bridge.settings import JsonSettings


def write_workbook(path: Path, *, status="completed", sn="SN-001", legacy=False, rows=None, extra_spectrum=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试数据"
    sheet.append([None, "测试站别", path.parent.name, "测试温度", "水冷25°C"])
    sheet.append(["电流(A)", "电压(V)", "功率(W)", "电光效率", "中心波长(nm)", "质心波长(nm)", "FWHM(nm)", "PIB", "SMSR(dB)"])
    for row in rows if rows is not None else [[2, 10, 0, 0, 976, 975.9, 0.4, .95], [4, 12, 24, .5, 977, 976.9, .5, .96]]:
        sheet.append(row)
    if extra_spectrum:
        sheet.cell(2, 10, "2 A")
        for row in range(3, 503):
            sheet.cell(row, 10, 970 + row / 100)
            sheet.cell(row, 11, row)
    if not legacy:
        session = workbook.create_sheet("会话信息")
        for key, value in {"SN": sn, "测试站别": path.parent.name, "测试温度": "水冷25°C",
                           "流程状态": status, "开始时间 UTC": "2026-09-10T01:00:00+00:00",
                           "出货报告格式版本": "1", "测试模式": "automatic", "会话编号": path.stem}.items():
            session.append([key, value])
        session.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


class DataBrowserTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        browser.clear_data_browser_cache()

    def list(self, **params):
        return browser.list_data_records({"rootPath": str(self.root), **params})

    def test_metadata_hidden_sheets_scalar_points_and_original_unchanged(self):
        path = self.root / "SN-001" / "封盖前" / "one.xlsx"
        write_workbook(path, extra_spectrum=True)
        before = hashlib.sha256(path.read_bytes()).digest()
        detail = browser.inspect_data_record({"path": str(path)})
        self.assertEqual(detail["record"]["temperature"], "水冷25°C")
        self.assertEqual(detail["record"]["timeSource"], "session")
        self.assertEqual(len(detail["points"]), 2)
        self.assertEqual(detail["points"][1]["efficiencyPercent"], 50)
        self.assertEqual(detail["points"][0]["pibPercent"], 95)
        self.assertIsNone(detail["points"][0]["smsrDb"])
        self.assertNotIn("wavelength", detail["points"][0])
        self.assertNotIn("spectra", detail)
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), before)

        # When includeSpectra is requested, spectra must be parsed and returned
        with_spectra = browser.inspect_data_record({"path": str(path), "includeSpectra": True})
        self.assertIn("spectra", with_spectra)
        self.assertIsNotNone(with_spectra["spectra"])
        self.assertEqual(len(with_spectra["spectra"]), 1)
        spec = with_spectra["spectra"][0]
        self.assertEqual(spec["currentA"], 2.0)
        self.assertEqual(len(spec["wavelengthNm"]), 500)
        self.assertEqual(len(spec["intensity"]), 500)
        self.assertAlmostEqual(spec["wavelengthNm"][0], 970.03)
        self.assertEqual(spec["intensity"][0], 3)
        self.assertEqual(spec["maxIntensity"], 502)


    def test_all_records_include_interrupted_legacy_broken_but_not_lockfiles(self):
        write_workbook(self.root / "SN-001" / "前" / "a.xlsx")
        write_workbook(self.root / "SN-001" / "后" / "b.xlsx", status="aborted_safely")
        write_workbook(self.root / "SN-001" / "后" / "c.xlsx", legacy=True)
        (self.root / "broken.xlsx").write_text("not a zip")
        (self.root / "~$temporary.xlsx").write_text("locked")
        result = self.list()
        self.assertEqual(result["total"], 4)
        self.assertEqual(result["errorCount"], 1)
        self.assertEqual(set(result["statuses"]), {"completed", "aborted_safely", "unknown"})
        legacy = next(r for r in result["records"] if r["fileName"] == "c.xlsx")
        self.assertEqual((legacy["sn"], legacy["snSource"]), ("SN-001", "directory"))
        self.assertEqual(legacy["timeSource"], "modified")

    def test_filters_pagination_and_dates(self):
        for index in range(4):
            write_workbook(self.root / "站别" / f"{index}.xlsx", sn=f"SN-{index}")
        result = self.list(sn="sn-", pageSize=2, page=2)
        self.assertEqual((result["total"], result["pageCount"], len(result["records"])), (4, 2, 2))
        self.assertEqual(self.list(sn="SN-2")["total"], 1)
        self.assertEqual(self.list(station="不存在")["total"], 0)
        self.assertEqual(self.list(status="aborted_safely")["total"], 0)
        stamp = result["records"][0]["timeMs"]
        self.assertEqual(self.list(fromMs=stamp, toMs=stamp + 1)["total"], 4)
        self.assertEqual(self.list(toMs=stamp)["total"], 0)

    def test_refresh_reuses_summary_even_after_detail_eviction_and_detects_changes(self):
        with patch.object(browser, "MAX_DETAIL_CACHE", 2):
            for index in range(5):
                write_workbook(self.root / "站别" / f"{index}.xlsx")
            with patch.object(browser, "load_workbook", wraps=browser.load_workbook) as load:
                self.list()
                self.list(refresh=True)
                self.assertEqual(load.call_count, 5)
                self.assertEqual(len(browser._DETAILS), 2)
                path = self.root / "站别" / "0.xlsx"
                stat = path.stat()
                os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 1000000000))
                self.list(refresh=True)
                self.assertEqual(load.call_count, 6)
                path.unlink()
                self.assertEqual(self.list(refresh=True)["total"], 4)
                with self.assertRaisesRegex(ValueError, "移动或删除"):
                    browser.inspect_data_record({"path": str(path)})

    def test_repaired_and_unsupported_workbooks(self):
        path = self.root / "fixed.xlsx"
        path.write_text("broken")
        self.assertEqual(self.list()["errorCount"], 1)
        write_workbook(path)
        self.assertEqual(self.list(refresh=True)["errorCount"], 0)
        workbook = Workbook(); workbook.active.append(["unrelated", "spreadsheet"])
        workbook.save(self.root / "other.xlsx"); workbook.close()
        self.assertEqual(self.list(refresh=True)["errorCount"], 1)

    def test_index_and_detail_describe_a_broken_file_the_same_way(self):
        path = self.root / "broken.xlsx"
        path.write_text("not a zip")
        record = next(r for r in self.list()["records"] if r["fileName"] == "broken.xlsx")
        self.assertEqual(record["readStatus"], "error")
        self.assertTrue(record["error"].startswith("无法读取 Excel："), record["error"])
        self.assertNotEqual(record["error"], "File is not a zip file")
        with self.assertRaises(ValueError) as raised:
            browser.inspect_data_record({"path": str(path)})
        self.assertEqual(record["error"], str(raised.exception))

    def test_duplicate_currents_single_point_missing_values_are_preserved(self):
        path = self.root / "one.xlsx"
        write_workbook(path, rows=[[2, None, 5], [2, None, 6], [4, 10, None]])
        detail = browser.inspect_data_record({"path": str(path)})
        self.assertEqual(detail["record"]["validPointCount"], 2)
        self.assertEqual(len(detail["points"]), 3)
        self.assertEqual([p["rowId"] for p in detail["points"]], [3, 4, 5])
        self.assertIsNone(detail["points"][2]["powerW"])
        write_workbook(self.root / "single.xlsx", rows=[[2, 10, 3]])
        self.assertEqual(self.list()["total"], 2)

    def test_explicit_percent_header_is_not_multiplied(self):
        path = self.root / "percent.xlsx"
        workbook = Workbook()
        workbook.active.append(["Current (A)", "Power (W)", "Efficiency (%)", "PIB (%)"])
        workbook.active.append([2, 1, .5, 90])
        workbook.save(path); workbook.close()
        detail = browser.inspect_data_record({"path": str(path)})
        self.assertEqual(detail["points"][0]["efficiencyPercent"], .5)
        self.assertEqual(detail["points"][0]["pibPercent"], 90)

    def test_parameters_and_cache_results_are_isolated(self):
        write_workbook(self.root / "one.xlsx")
        for params in ({"refresh": "false"}, {"page": True}, {"pageSize": 999}, {"fromMs": "NaN"}, {"fromMs": 2, "toMs": 1}):
            with self.assertRaises(ValueError): self.list(**params)
        result = self.list(); result["records"][0]["sn"] = "mutated"
        self.assertEqual(self.list()["records"][0]["sn"], "SN-001")
        with self.assertRaisesRegex(ValueError, "目录不存在"):
            browser.list_data_records({"rootPath": str(self.root / "missing")})

    def test_file_task_bridge_and_light_snapshot(self):
        path = self.root / "one.xlsx"; write_workbook(path)
        backend = HeadlessBackend(JsonSettings(self.root / "settings.json"))
        service = BridgeService(backend)
        for method, params in (("dataBrowser.list", {"rootPath": str(self.root)}), ("dataBrowser.inspect", {"path": str(path)})):
            self.assertIn(method, BACKGROUND_FILE_TASK_METHODS)
            response = service.handle_line(json.dumps({"v": 2, "id": method, "method": method, "params": params}), file_task=True)
            self.assertNotIn("error", response)
            self.assertIn("result", response)
        snapshot = backend.snapshot({"view": "data"})
        self.assertEqual(snapshot["measurements"], {})
        self.assertNotIn("points", snapshot["pd"])


if __name__ == "__main__":
    unittest.main()
