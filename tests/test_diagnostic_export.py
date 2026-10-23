"""Real ZIP round trips and lifecycle checks, with no instrument access."""

import hashlib
import json
import tempfile
import threading
import unittest
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from openpyxl import load_workbook

from combined_test.automation import AutomaticTestState
from combined_test.excel_export import ExcelTestRecord, save_test_records
from combined_test.record_store import SessionStatus
from tauri_bridge.headless_backend import HeadlessBackend
from tauri_bridge.service import BACKGROUND_FILE_TASK_METHODS, BridgeService
from tauri_bridge.settings import JsonSettings


class DiagnosticExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.backend = HeadlessBackend(JsonSettings(self.root / "settings.json"))
        self.station = self.backend.runtime
        self.manager = self.station.power_stability
        self.samples = [(i * 0.1, (i % 2) * 0.15) for i in range(1801)]
        self.manager._finish_learning(self.manager.job_epoch, self.samples, self.manager.identity(), self.manager.policy())
        self.manager.freeze()
        self.frozen = self.manager.model
        self.assertIsNotNone(self.frozen)
        store = self.station.record_store
        self.workbook = store.begin_session(self.root / "Excel", "样品-A", datetime.now(), mode="automatic",
                                            settings={**self.station.configuration, "powerStability": self.manager.session_metadata()})
        record = ExcelTestRecord(1, 12, 32.5, 0.1, 976, 976, 0.4, 0.99, [975, 976], [1, 3])
        store.queue(record)
        store.record_stability_diagnostic(1, self.manager.snapshot())
        store.complete_session(SessionStatus.COMPLETED, "完成", shutdown_confirmed=True)
        save_test_records(self.workbook, (record,), session=store.current_session,
                          attempts=store.list_attempts(store.current_session.session_id))
        store.mark_saved((record,))
        self.station.automatic.state = AutomaticTestState.COMPLETED
        self.station.automatic_terminal_outcome = "succeeded"
        self.station.export_status = "completed"
        self.station.export_path = str(self.workbook)
        self.session = store.current_session
        self.journal = store.journal_path
        self.params = {"sessionId": self.session.session_id}

    def tearDown(self):
        self.backend.export_worker.graceful_stop(timeout_s=5)
        self.directory.cleanup()

    def export(self):
        return self.backend.dispatch_file_task("automatic.exportDiagnostics", self.params)

    def test_bundle_contains_exact_trace_journal_workbook_and_verified_inventory(self):
        self.backend.settings.set("unrelatedSecret", "must-not-export")
        originals = {path: path.read_bytes() for path in (self.workbook, self.journal, self.backend.settings.path)}
        result = self.export()
        self.assertEqual(result["sessionId"], self.session.session_id)
        self.assertEqual(result["warnings"], [])
        with ZipFile(result["path"]) as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual(archive.read("session.jsonl"), originals[self.journal])
            self.assertEqual(archive.read(f"result/{self.workbook.name}"), originals[self.workbook])
            bg = json.loads(archive.read("background.json"))
            self.assertEqual(bg["background"]["model_id"], self.frozen.model_id)
            self.assertEqual(bg["backgroundLearning"]["samples"], [list(item) for item in self.samples])
            self.assertNotIn("settings.json", archive.namelist())
            self.assertNotIn(b"must-not-export", b"".join(archive.read(name) for name in archive.namelist()))
            manifest = json.loads(archive.read("manifest.json"))
            self.assertTrue(manifest["complete"])
            self.assertEqual(manifest["sessionId"], self.session.session_id)
            for item in manifest["files"]:
                data = archive.read(item["name"])
                self.assertEqual(hashlib.sha256(data).hexdigest(), item["sha256"])
                self.assertEqual(len(data), item["size"])
        for path, original in originals.items():
            self.assertEqual(path.read_bytes(), original)
        second = self.export()
        self.assertNotEqual(second["path"], result["path"])
        self.assertTrue(Path(result["path"]).is_file())

    def test_relearning_after_the_run_never_replaces_the_frozen_learning_trace(self):
        old_id = self.frozen.model_id
        self.manager._finish_learning(self.manager.job_epoch, [(t, p * 2) for t, p in self.samples],
                                       self.manager.identity(), self.manager.policy())
        self.assertNotEqual(self.manager.model.model_id, old_id)
        with ZipFile(self.export()["path"]) as archive:
            bg = json.loads(archive.read("background.json"))
            self.assertEqual(bg["backgroundLearning"]["model"]["model_id"], old_id)
            self.assertEqual(bg["backgroundLearning"]["samples"], [list(item) for item in self.samples])

    def test_full_baseline_is_retained_without_truncating_excel_model_metadata(self):
        metadata = self.session.settings["powerStability"]
        self.assertGreater(len(str(metadata)), 32767)
        workbook = load_workbook(self.workbook, read_only=True, data_only=True)
        try:
            rows = dict(workbook["会话信息"].iter_rows(min_col=1, max_col=2, values_only=True))
            value = rows["powerStability"]
            self.assertIn(self.frozen.model_id, value)
            self.assertNotIn("backgroundLearning", value)
            self.assertLess(len(value), 32767)
        finally:
            workbook.close()
        self.assertEqual(len(metadata["backgroundLearning"]["samples"]), 1801)

    def test_old_session_trace_lookup_requires_exact_model_id(self):
        settings = dict(self.session.settings)
        settings["powerStability"] = {**settings["powerStability"], "backgroundLearning": None}
        self.station.record_store.current_session = replace(self.session, settings=settings)
        self.assertEqual(self.export()["warnings"], [])
        saved = self.backend.settings.get("runtime/powerBackgrounds")
        saved[self.frozen.identity]["model"]["model_id"] = "different-model"
        self.backend.settings.set("runtime/powerBackgrounds", saved)
        result = self.export()
        self.assertTrue(any("原始样本缺失" in warning for warning in result["warnings"]))
        with ZipFile(result["path"]) as archive:
            self.assertIsNone(json.loads(archive.read("background.json"))["backgroundLearning"])
        saved[self.frozen.identity]["model"]["model_id"] = self.frozen.model_id
        del saved[self.frozen.identity]["samples"]
        self.backend.settings.set("runtime/powerBackgrounds", saved)
        result = self.export()
        self.assertTrue(any("原始样本缺失" in warning for warning in result["warnings"]))
        with ZipFile(result["path"]) as archive:
            self.assertFalse(json.loads(archive.read("manifest.json"))["complete"])

    def test_failed_excel_export_still_provides_journal_background_and_leftover_csv(self):
        self.station.export_status = "failed"
        self.station.export_error = "Excel 被占用"
        csv = self.workbook.with_suffix(".csv")
        csv.write_text("elapsed_s,power_w\n0,32.5\n", encoding="utf-8")
        self.station.automatic_sample_path = str(csv)
        result = self.export()
        self.assertTrue(any("Excel 被占用" in warning for warning in result["warnings"]))
        with ZipFile(result["path"]) as archive:
            self.assertNotIn(f"result/{self.workbook.name}", archive.namelist())
            self.assertEqual(archive.read(f"result/{csv.name}"), csv.read_bytes())
            self.assertFalse(json.loads(archive.read("manifest.json"))["complete"])

    def test_wrong_round_active_run_or_pending_writes_do_not_export(self):
        for params in ({}, {"sessionId": "old-session"}, {"sessionId": None}):
            with self.assertRaisesRegex(ValueError, "轮次"):
                self.backend.dispatch_file_task("automatic.exportDiagnostics", params)
        for state, status in ((AutomaticTestState.WAITING_STABLE, "completed"),
                              (AutomaticTestState.COMPLETED, "exporting")):
            self.station.automatic.state, self.station.export_status = state, status
            with self.assertRaisesRegex(ValueError, "写入完成"):
                self.export()
        self.assertFalse((self.root / "diagnostics").exists())

    def test_mismatched_or_missing_journal_fails_without_a_partial_package(self):
        self.journal.write_text(json.dumps({"type": "session_start", "session": {"session_id": "other"}}))
        with self.assertRaisesRegex(ValueError, "不匹配"):
            self.export()
        self.journal.unlink()
        with self.assertRaisesRegex(ValueError, "日志不存在"):
            self.export()
        self.assertFalse((self.root / "diagnostics").exists())

    def test_zip_write_failure_cleans_temporary_file_and_allows_retry(self):
        with patch("combined_test.diagnostics.os.replace", side_effect=OSError("disk full")):
            with self.assertRaisesRegex(OSError, "disk full"):
                self.export()
        self.assertEqual(list((self.root / "diagnostics").iterdir()), [])
        self.assertTrue(Path(self.export()["path"]).is_file())

    def test_source_change_during_compression_never_publishes_a_mixed_package(self):
        original_open = ZipFile.open

        def changing_open(archive, name, mode="r", *args, **kwargs):
            if name == f"result/{self.workbook.name}" and mode == "w":
                with self.workbook.open("ab") as writer:
                    writer.write(b"changed while exporting")
            return original_open(archive, name, mode, *args, **kwargs)

        with patch.object(ZipFile, "open", changing_open):
            with self.assertRaisesRegex(ValueError, "仍在写入"):
                self.export()
        self.assertEqual(list((self.root / "diagnostics").iterdir()), [])

    def test_compression_does_not_hold_station_lock_and_duplicate_requests_are_blocked(self):
        entered, release = threading.Event(), threading.Event()
        result = []

        def compress(**kwargs):
            entered.set()
            self.assertTrue(release.wait(2))
            return {"path": "test.zip", "sessionId": self.session.session_id, "warnings": []}

        with patch("tauri_bridge.diagnostic_export.export_diagnostic_bundle", side_effect=compress):
            thread = threading.Thread(target=lambda: result.append(self.export()))
            thread.start()
            try:
                self.assertTrue(entered.wait(2))
                self.assertTrue(self.station._lock.acquire(blocking=False))
                self.station._lock.release()
                self.backend.snapshot()
                with self.assertRaisesRegex(ValueError, "正在导出"):
                    self.export()
            finally:
                release.set()
                thread.join(3)
        self.assertEqual(len(result), 1)

    def test_file_worker_bridge_routes_export_and_rejects_stale_station(self):
        self.assertIn("automatic.exportDiagnostics", BACKGROUND_FILE_TASK_METHODS)
        response = BridgeService(self.backend).handle_line(json.dumps({"v": 2, "id": "diagnostic", "method": "automatic.exportDiagnostics", "params": self.params}), file_task=True)
        self.assertEqual(response["result"]["sessionId"], self.session.session_id)
        with self.assertRaisesRegex(ValueError, "单工位"):
            self.backend.dispatch_file_task("automatic.exportDiagnostics", {**self.params, "stationId": "2"})


if __name__ == "__main__":
    unittest.main()
